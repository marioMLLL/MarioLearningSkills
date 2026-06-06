#!/usr/bin/env python3
"""
Extract text from PDF, DOCX, or DOC files.
Usage: python3 extract_text.py <file_path>
Outputs extracted text to stdout.
"""

import sys
import os
import subprocess


def extract_pdf(path):
    """Extract text from PDF using available libraries."""
    # Try PyMuPDF (fitz) first - best quality
    try:
        import fitz
        doc = fitz.open(path)
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text
    except ImportError:
        pass

    # Try PyPDF2
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    except ImportError:
        pass

    # Try pdfplumber
    try:
        import pdfplumber
        text = ""
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""
        return text
    except ImportError:
        pass

    # Fallback: use pdftotext command line tool
    try:
        result = subprocess.run(
            ["pdftotext", "-layout", path, "-"],
            capture_output=True, text=True
        )
        if result.returncode == 0:
            return result.stdout
    except FileNotFoundError:
        pass

    raise RuntimeError(
        "No PDF extraction library found. Install one of: "
        "PyMuPDF (fitz), PyPDF2, pdfplumber, or poppler-utils (pdftotext)"
    )


def extract_docx(path):
    """Extract text from DOCX file."""
    try:
        from docx import Document
        doc = Document(path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        # Also extract table content
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        return text
    except ImportError:
        raise RuntimeError(
            "python-docx not installed. Run: pip install python-docx"
        )


def extract_doc(path):
    """Extract text from legacy DOC file."""
    # Try using antiword
    try:
        result = subprocess.run(
            ["antiword", path],
            capture_output=True, text=True
        )
        if result.returncode == 0:
            return result.stdout
    except FileNotFoundError:
        pass

    # Try using libreoffice to convert to docx, then extract
    try:
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            result = subprocess.run(
                [
                    "libreoffice", "--headless", "--convert-to", "docx",
                    "--outdir", tmpdir, path
                ],
                capture_output=True, text=True, timeout=60
            )
            if result.returncode == 0:
                docx_file = os.path.join(
                    tmpdir,
                    os.path.splitext(os.path.basename(path))[0] + ".docx"
                )
                if os.path.exists(docx_file):
                    return extract_docx(docx_file)
    except FileNotFoundError:
        pass
    except subprocess.TimeoutExpired:
        pass

    raise RuntimeError(
        "Cannot extract DOC file. Install antiword or libreoffice: "
        "brew install antiword  or  brew install --cask libreoffice"
    )


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 extract_text.py <file_path>", file=sys.stderr)
        sys.exit(1)

    file_path = sys.argv[1]
    if not os.path.exists(file_path):
        print(f"Error: File not found: {file_path}", file=sys.stderr)
        sys.exit(1)

    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".pdf":
            text = extract_pdf(file_path)
        elif ext == ".docx":
            text = extract_docx(file_path)
        elif ext == ".doc":
            text = extract_doc(file_path)
        else:
            print(f"Error: Unsupported file type: {ext}", file=sys.stderr)
            sys.exit(1)

        if not text.strip():
            print("Warning: Extracted text is empty.", file=sys.stderr)
            sys.exit(1)

        print(text)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
